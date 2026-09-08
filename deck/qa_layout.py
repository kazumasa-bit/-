"""Analytical layout QA: text overflow, slide bounds, text-on-text overlap."""
import sys, math, unicodedata
from pptx import Presentation
from pptx.util import Emu

EMU_IN = 914400.0
SLIDE_W, SLIDE_H = 13.333, 7.5

def adv(ch, pt):
    if ch == '\t': return pt * 2
    w = unicodedata.east_asian_width(ch)
    return pt * (1.0 if w in ('W', 'F') else 0.52)

def para_text(p): return ''.join(r.text for r in p.runs)

def para_size(p, default=18.0):
    for r in p.runs:
        if r.font.size: return r.font.size.pt
    if p.font.size: return p.font.size.pt
    return default

def check(path):
    prs = Presentation(path)
    issues = []
    for si, slide in enumerate(prs.slides, 1):
        boxes = []
        for sh in slide.shapes:
            if sh.left is None: continue
            x, y = sh.left / EMU_IN, sh.top / EMU_IN
            w, h = (sh.width or 0) / EMU_IN, (sh.height or 0) / EMU_IN
            # slide bounds
            if x < -0.02 or y < -0.02 or x + w > SLIDE_W + 0.02 or y + h > SLIDE_H + 0.02:
                issues.append((si, 'BOUNDS', f'{sh.shape_type} x={x:.2f} y={y:.2f} w={w:.2f} h={h:.2f}'))
            if not sh.has_text_frame: continue
            tf = sh.text_frame
            txt = tf.text
            if not txt.strip(): continue
            li = (tf.margin_left or 0) / EMU_IN
            ri = (tf.margin_right or 0) / EMU_IN
            ti = (tf.margin_top or 0) / EMU_IN
            bi = (tf.margin_bottom or 0) / EMU_IN
            usable_w = max(w - li - ri, 0.15)
            need_h = 0.0
            wrapped = 0
            for p in tf.paragraphs:
                t = para_text(p)
                pt = para_size(p)
                ls = p.line_spacing
                if ls is None:
                    line_h = pt * 1.22 / 72.0
                elif hasattr(ls, 'pt'):        # Length -> exact points
                    line_h = ls.pt / 72.0
                else:                          # float -> multiple of line
                    line_h = pt * float(ls) * 1.2 / 72.0
                sa = (p.space_after.pt / 72.0) if p.space_after else 0.0
                if not t:
                    need_h += line_h; continue
                bullet_pad = 0.22 if '<a:buChar' in p._p.xml or '<a:buAutoNum' in p._p.xml else 0.0
                width_pt = sum(adv(c, pt) for c in t)
                lines = max(1, math.ceil((width_pt / 72.0) / max(usable_w - bullet_pad, 0.15)))
                wrapped += lines - 1
                need_h += lines * line_h + sa
            avail_h = h - ti - bi
            if need_h > avail_h * 1.06 and avail_h > 0:
                issues.append((si, 'OVERFLOW',
                               f'"{txt[:34].replace(chr(10)," / ")}" need={need_h:.2f}in avail={avail_h:.2f}in '
                               f'(box {w:.2f}x{h:.2f} @ {x:.2f},{y:.2f})'))
            boxes.append((x + li, y + ti, w - li - ri, min(need_h, h), txt))
        # text-on-text overlap
        for i in range(len(boxes)):
            for j in range(i + 1, len(boxes)):
                a, b = boxes[i], boxes[j]
                ox = min(a[0] + a[2], b[0] + b[2]) - max(a[0], b[0])
                oy = min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])
                if ox > 0.06 and oy > 0.06:
                    area = ox * oy
                    if area > 0.05:
                        issues.append((si, 'TEXT-OVERLAP',
                                       f'"{a[4][:20]}" x "{b[4][:20]}" overlap {ox:.2f}x{oy:.2f}in'))
    return issues

iss = check(sys.argv[1])
if not iss:
    print('LAYOUT OK — no overflow, bounds, or text-overlap issues found.')
for s, k, d in iss:
    print(f'slide {s:>2}  {k:<13} {d}')
print(f'\ntotal issues: {len(iss)}')
